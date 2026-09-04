# -*- coding: utf-8 -*-
"""把申请表 §三「主要功能和技术特点」拆成两段：主要功能(>=500字) + 技术特点(<=100字)。"""
import os, re
from docx import Document
from docx.shared import RGBColor
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph

F1 = r"D:\workspace\tiktok\mini-game\docs\软著\1-软件著作权登记申请表填写指南.docx"
BLACK = RGBColor(0, 0, 0)

FUNC = (
    "本软件是一款运行于抖音小游戏平台的休闲益智类合成游戏，采用 Cocos Creator 引擎与 "
    "TypeScript 语言开发，竖屏 720×1280 分辨率适配，包含加载、主界面、游戏、排行榜、设置五个"
    "场景。软件启动后进入加载场景，加载并显示游戏资源的加载进度，完成后自动进入主界面；主界面"
    "提供开始游戏、排行榜、设置三个入口，并在当日首次进入时弹出每日礼包窗口，供玩家领取猫币奖"
    "励，亦可观看激励视频使奖励翻倍。进入游戏后，玩家通过在屏幕上按住并左右拖动来控制待投放甜"
    "品的水平位置，松开手指后甜品落入容器，容器内两个相同等级的甜品发生碰撞时自动合成为更高一"
    "级的甜品，并获得相应分数。游戏场景顶部最多同时排布三位猫咪顾客，每位顾客通过头顶气泡展示"
    "其所需甜品的等级与数量，玩家合成出对应等级的甜品即可为其满足一份订单；当一位顾客的全部订"
    "单被满足后该顾客离场并由后续顾客补位，满足本关全部顾客订单即判定通关，弹出胜利结算弹窗，"
    "依据本局得分给出一至三星评定并发放猫币奖励；若容器内甜品堆积超出顶部警戒线并持续一定时"
    "间，则本局失败，弹出失败结算弹窗。游戏过程中还提供锤子、洗牌等道具及观看激励视频获取金币"
    "的功能，并支持在胜利结算时看广告使奖励翻倍、在失败结算时看广告复活继续本局。此外，软件提"
    "供暂停功能、排行榜功能（通过网络接口提交并查询玩家与好友的成绩排名）、背景音乐与音效的开"
    "关设置功能，以及将战绩分享炫耀的功能。游戏进度、设置项与每日礼包领取状态等数据保存在用户"
    "设备本地，玩家成绩则通过网络接口上传至服务端，用于排行榜展示与跨设备进度同步。"
)

TECH = (
    "采用场景—组件化架构，游戏逻辑（合成判定、溢出检测、计分、顾客订单匹配）与界面表现分离；甜"
    "品属性与关卡顾客需求配置化驱动，便于扩展；接入抖音平台接口实现安全区适配，并针对移动端优"
    "化资源加载与渲染性能。"
)

def cjk(s):
    return len(re.findall(r'[一-鿿]', s))

print("主要功能 汉字数 =", cjk(FUNC), "(要求 >=500)")
print("技术特点 汉字数 =", cjk(TECH), "(要求 <=100)")
assert cjk(FUNC) >= 500, "主要功能不足 500 字"
assert cjk(TECH) <= 100, "技术特点超过 100 字"

d = Document(F1)

# 找到原 §三 正文段（之前合并写入、以「本软件是一款」开头）
body = None
for p in d.paragraphs:
    if p.text.strip().startswith("本软件是一款运行于抖音小游戏平台"):
        body = p
        break
assert body is not None, "未找到 §三 正文段"

ref = body.runs[0].font
ref_size = ref.size
ref_name = ref.name

def style_run(r):
    r.font.color.rgb = BLACK
    if ref_size:
        r.font.size = ref_size
    if ref_name:
        r.font.name = ref_name
    rpr = r._element.get_or_add_rPr()
    from docx.oxml.ns import qn
    rf = OxmlElement('w:rFonts')
    rf.set(qn('w:eastAsia'), "宋体")
    rpr.append(rf)
    return r

def fill_para(p, label, content):
    # 清空原 run
    for run in list(p.runs):
        run._element.getparent().remove(run._element)
    lab = p.add_run(label)
    lab.font.bold = True
    style_run(lab)
    txt = p.add_run(content)
    txt.font.bold = False
    style_run(txt)

def insert_after(paragraph):
    new_p = OxmlElement('w:p')
    paragraph._p.addnext(new_p)
    return Paragraph(new_p, paragraph._parent)

# 段一：主要功能
fill_para(body, "主要功能：", FUNC)
# 段二：技术特点（插在主要功能之后）
tech_p = insert_after(body)
fill_para(tech_p, "技术特点：", TECH)

d.save(F1)
print("SAVED:", F1)
