# -*- coding: utf-8 -*-
"""
生成软著「源代码鉴别材料」成稿 .docx。
- 按模板顺序拼接 35 个自有 .ts 文件（排除 tt.d.ts 与 data/LevelGenerator.ts）
- 文件间加 // ===== 路径 ===== 分隔
- 9 磅等宽字体、单倍行距、A4 窄边距
- 每 50 逻辑行强制分页；超 60 页则取前 30 页 + 后 30 页（一般交存）
"""
import os
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_BREAK
from docx.enum.section import WD_SECTION
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

ROOT = r"D:\workspace\tiktok\mini-game\client\assets\scenes\scripts"
OUT = r"D:\workspace\tiktok\mini-game\docs\软著\3-源代码鉴别材料-成稿.docx"
SOFT_NAME = "猫咪甜品店软件 V1.0"
LINES_PER_PAGE = 50
MAX_PAGES_EACH = 30  # 前后各 30 页

# 模板拼接顺序（35 个文件，相对 ROOT）
ORDER = [
    "data/GameTypes.ts", "data/GameState.ts", "data/DessertConfig.ts",
    "core/Container.ts", "core/Dessert.ts", "core/DropController.ts",
    "core/MergeManager.ts", "core/OverflowDetector.ts", "core/ScoreManager.ts",
    "core/ItemManager.ts", "core/CustomerManager.ts",
    "ui/LoadingScene.ts", "ui/HomeScene.ts", "ui/GameScene.ts", "ui/RankScene.ts",
    "ui/SettingsScene.ts", "ui/DesignTokens.ts", "ui/GlobalFontManager.ts",
    "ui/PopupManager.ts",
    "ui/popups/PopupUIHelper.ts", "ui/popups/WinPopup.ts", "ui/popups/FailPopup.ts",
    "ui/popups/PausePopup.ts", "ui/popups/DailyGiftPopup.ts", "ui/popups/RankPopup.ts",
    "ui/popups/SettingsPopup.ts",
    "platform/DouyinSDK.ts", "platform/SafeArea.ts", "platform/AdConfig.ts",
    "net/ApiConfig.ts", "net/ApiTypes.ts", "net/ApiClient.ts",
    "utils/AudioManager.ts", "utils/Toast.ts", "utils/TouchSpace.ts",
]

# --- 组装行流 ---
code_lines = 0          # 纯源代码行（用于源程序量统计）
display = []             # 文档实际呈现的行（含分隔注释）
for rel in ORDER:
    path = os.path.join(ROOT, *rel.split("/"))
    with open(path, "r", encoding="utf-8") as f:
        lines = f.read().split("\n")
    if lines and lines[-1] == "":
        lines = lines[:-1]  # 去掉文件末尾空行
    display.append(f"// ===== {rel} =====")
    display.extend(lines)
    code_lines += len(lines)

total_disp = len(display)
print(f"35 文件源代码行数(源程序量) = {code_lines}")
print(f"含分隔注释总呈现行 = {total_disp}")

cap = LINES_PER_PAGE * MAX_PAGES_EACH  # 1500
if total_disp <= cap * 2:
    pages_lines = display
    trimmed = False
else:
    pages_lines = display[:cap] + display[-cap:]
    trimmed = True
print(f"输出行数 = {len(pages_lines)}  裁切={trimmed}")

# --- 构建 docx ---
doc = Document()
sec = doc.sections[0]
sec.page_height = Cm(29.7)
sec.page_width = Cm(21.0)
sec.top_margin = Cm(1.4)
sec.bottom_margin = Cm(1.4)
sec.left_margin = Cm(1.5)
sec.right_margin = Cm(1.5)

# 页眉：软件全称（三处需逐字一致）
hdr = sec.header.paragraphs[0]
hdr.text = SOFT_NAME + "　源程序"
for r in hdr.runs:
    r.font.size = Pt(9)

def make_line(text, page_break=False):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.line_spacing = 1.0
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    r = p.add_run()
    if page_break:
        r.add_break(WD_BREAK.PAGE)
    r2 = p.add_run(text if text != "" else "")
    for run in (r, r2):
        run.font.size = Pt(9)
        run.font.name = "Consolas"
        rpr = run._element.get_or_add_rPr()
        rfonts = rpr.find(qn('w:rFonts'))
        if rfonts is None:
            rfonts = OxmlElement('w:rFonts')
            rpr.append(rfonts)
        rfonts.set(qn('w:ascii'), "Consolas")
        rfonts.set(qn('w:hAnsi'), "Consolas")
        rfonts.set(qn('w:eastAsia'), "宋体")
    return p

for i, line in enumerate(pages_lines):
    brk = (i != 0 and i % LINES_PER_PAGE == 0)
    if trimmed and i == cap:
        brk = True  # 前后衔接处分页
    make_line(line, page_break=brk)

doc.save(OUT)
print("SAVED:", OUT)
