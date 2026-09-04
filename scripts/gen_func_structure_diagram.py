# -*- coding: utf-8 -*-
"""
生成「软件功能结构图」（树状）并插入到 docs/软著/2-操作说明书模板.docx 的
「插图 2-1」占位处，占位段落转为图注「图 2-1 软件功能结构图」。

数据来源：操作说明书第 2 节的功能模块表（8 个模块 + 关键子功能）。
依赖：Pillow、python-docx（仓库已装）。中文字体优先用微软雅黑，回落项目原字。

用法：python scripts/gen_func_structure_diagram.py
"""
import os
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DOCX = os.path.join(ROOT, "docs", "软著", "2-操作说明书模板.docx")
PNG = os.path.join(ROOT, "docs", "软著", "_func_structure_2-1.png")

SOFT_NAME = "一起开猫店"

# 8 个模块 + 关键子功能（取自功能模块表，凝练为短标签）
MODULES = [
    ("加载模块",     ["资源加载", "加载进度", "进入主界面"]),
    ("主界面模块",   ["开始游戏", "排行榜入口", "设置入口", "每日礼包"]),
    ("游戏核心模块", ["甜品投放", "碰撞合成", "溢出检测", "实时计分"]),
    ("排行榜模块",   ["成绩提交", "排名展示"]),
    ("设置模块",     ["音乐开关", "音效开关"]),
    ("弹窗系统",     ["胜利/失败结算", "暂停弹窗", "每日礼包"]),
    ("平台接入模块", ["接口调用", "安全区适配"]),
    ("数据存储模块", ["本地存储", "数据读取"]),
]

# 配色（暖色，呼应游戏但保持文档严肃）
C_BG = (255, 255, 255)
C_LINE = (120, 100, 84)
C_ROOT_FILL = (240, 180, 120)
C_ROOT_BORDER = (150, 100, 60)
C_MOD_FILL = (255, 244, 224)
C_MOD_BORDER = (150, 110, 70)
C_SUB_FILL = (245, 245, 240)
C_SUB_BORDER = (165, 150, 130)
C_TEXT = (40, 30, 20)


def load_font(size):
    for path in (r"C:\Windows\Fonts\msyh.ttc", r"C:\Windows\Fonts\msyhbd.ttc",
                 os.path.join(ROOT, "backup", "font-orig-2026-06-11", "ZCOOLKuaiLe-Regular.ttf.orig")):
        if os.path.exists(path):
            try:
                return ImageFont.truetype(path, size, index=0)
            except Exception:
                continue
    return ImageFont.load_default()


F_ROOT = load_font(40)
F_MOD = load_font(30)
F_SUB = load_font(24)

# 布局参数
MARGIN = 70
MW = 300          # 模块列宽
GAP = 44          # 模块间距
MOD_H = 86
SUB_W = 244
SUB_H = 56
SUB_GAP = 18
ROOT_W, ROOT_H = 400, 100
Y_ROOT = 60
Y_MOD = 300
Y_SUB0 = 470

N = len(MODULES)
CANVAS_W = MARGIN * 2 + N * MW + (N - 1) * GAP
max_subs = max(len(s) for _, s in MODULES)
CANVAS_H = Y_SUB0 + max_subs * (SUB_H + SUB_GAP) + 40


def text_centered(draw, cx, cy, text, font, color=C_TEXT):
    l, t, r, b = draw.textbbox((0, 0), text, font=font)
    draw.text((cx - (r - l) / 2, cy - (b - t) / 2 - t), text, font=font, fill=color)


def rounded_box(draw, cx, cy, w, h, fill, border, radius=14, width=3):
    x0, y0, x1, y1 = cx - w / 2, cy - h / 2, cx + w / 2, cy + h / 2
    draw.rounded_rectangle([x0, y0, x1, y1], radius=radius, fill=fill, outline=border, width=width)


def build_png():
    img = Image.new("RGB", (CANVAS_W, CANVAS_H), C_BG)
    d = ImageDraw.Draw(img)

    root_cx = CANVAS_W / 2
    root_cy = Y_ROOT + ROOT_H / 2
    mod_cxs = [MARGIN + MW / 2 + i * (MW + GAP) for i in range(N)]
    mod_cy = Y_MOD + MOD_H / 2

    # 连接线：根 → 横向母线 → 各模块顶部
    bus_y = (root_cy + ROOT_H / 2 + (mod_cy - MOD_H / 2)) / 2
    d.line([(root_cx, root_cy + ROOT_H / 2), (root_cx, bus_y)], fill=C_LINE, width=3)
    d.line([(mod_cxs[0], bus_y), (mod_cxs[-1], bus_y)], fill=C_LINE, width=3)
    for cx in mod_cxs:
        d.line([(cx, bus_y), (cx, mod_cy - MOD_H / 2)], fill=C_LINE, width=3)

    # 模块 + 子功能
    for (name, subs), cx in zip(MODULES, mod_cxs):
        rounded_box(d, cx, mod_cy, MW - 24, MOD_H, C_MOD_FILL, C_MOD_BORDER)
        text_centered(d, cx, mod_cy, name, F_MOD)
        # 模块底部向下的主干，连到每个子功能
        prev_y = mod_cy + MOD_H / 2
        for j, sub in enumerate(subs):
            sy = Y_SUB0 + j * (SUB_H + SUB_GAP) + SUB_H / 2
            d.line([(cx, prev_y), (cx, sy - SUB_H / 2)], fill=C_LINE, width=2)
            rounded_box(d, cx, sy, SUB_W, SUB_H, C_SUB_FILL, C_SUB_BORDER, radius=12, width=2)
            text_centered(d, cx, sy, sub, F_SUB)
            prev_y = sy + SUB_H / 2

    # 根节点（最后画，盖住线头）
    rounded_box(d, root_cx, root_cy, ROOT_W, ROOT_H, C_ROOT_FILL, C_ROOT_BORDER, radius=18, width=4)
    text_centered(d, root_cx, root_cy, SOFT_NAME, F_ROOT, color=(60, 40, 20))

    img.save(PNG, "PNG")
    print(f"[diagram] PNG 生成: {os.path.relpath(PNG, ROOT)}  ({CANVAS_W}x{CANVAS_H})")


def insert_into_docx():
    doc = Document(DOCX)
    ph = None
    for p in doc.paragraphs:
        if "插图 2-1" in p.text or "插图2-1" in p.text:
            ph = p
            break
    if ph is None:
        raise SystemExit("未找到「插图 2-1」占位段落")

    # 图片段落：新建后移动到占位段落之前
    img_para = doc.add_paragraph()
    img_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    img_para.add_run().add_picture(PNG, width=Inches(6.3))
    ph._p.addprevious(img_para._p)

    # 占位段落转为图注
    for r in list(ph.runs):
        r._element.getparent().remove(r._element)
    run = ph.add_run("图 2-1 软件功能结构图")
    run.font.color.rgb = RGBColor(0, 0, 0)
    run.font.bold = False
    ph.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.save(DOCX)
    print(f"[diagram] 已插入并保存: {os.path.relpath(DOCX, ROOT)}")


if __name__ == "__main__":
    build_png()
    insert_into_docx()
